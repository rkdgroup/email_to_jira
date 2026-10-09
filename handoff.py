"""
EOD Hands-off Report
Lists DSLF tickets that moved into "Done" or "Waiting on Qty Approval" in the
last N hours, writes them to a dated Excel file under Documents/"Hands-off report/", and
optionally emails it through classic Outlook from the signed-in account.

A ticket that entered both statuses in the window shows once, with whichever
it entered last. Done rows come first (green), then Waiting rows (amber).
Re-running on the same day overwrites that day's file.

Usage:
    python handoff.py                     # last 8 hours, file only 
    python handoff.py --send              # + email to the .env recipients
    python handoff.py --hours 12 --send
    python handoff.py --send --to me@x.com    # test run: only to me, no CC

Recipients come from HANDOFF_EMAIL_TO / HANDOFF_EMAIL_CC in .env (comma-separated).
Passing --to replaces both, so a test send never copies anyone; add --cc to copy.
"""

import os
import re
import sys
import argparse
from datetime import datetime, timedelta, timezone
from pathlib import Path

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

load_dotenv(Path(__file__).parent / ".env")

sys.path.insert(0, str(Path(__file__).parent))
from tools_jira import _auth, _get_jira_base_url

PROJECT     = "DSLF"
DONE        = "Done"
WAITING     = "Waiting on Qty Approval"
STATUSES    = (DONE, WAITING)
MGR_ORDER   = "customfield_12192"  # Manager Order Number
REPORT_DIR  = Path.home() / "OneDrive" / "Documents" / "Hands-off report"
TITLE       = "EOD Hands-off Report"
EMAIL_BODY  = "<p>Hi,</p><p>Please find the EOD Hands-off report attached.</p>"

FILLS = {
    DONE:    PatternFill("solid", fgColor="C6EFCE"),  # green
    WAITING: PatternFill("solid", fgColor="FFEB9C"),  # amber
}


def _get(path: str, params: dict) -> dict:
    """GET from Jira; fail loudly so a broken query never reads as 'no tickets'."""
    resp = requests.get(f"{_get_jira_base_url()}{path}", auth=_auth(),
                        headers={"Accept": "application/json"},
                        params=params, timeout=30)
    if resp.status_code != 200:
        sys.exit(f"Jira request failed ({resp.status_code}): {resp.text[:300]}")
    return resp.json()


def find_issues(hours: int) -> list:
    clauses = " OR ".join(f'status CHANGED TO "{s}" AFTER "-{hours}h"' for s in STATUSES)
    jql = f"project = {PROJECT} AND ({clauses}) ORDER BY key"
    issues, token = [], None
    while True:
        params = {"jql": jql, "maxResults": 100, "fields": MGR_ORDER}
        if token:
            params["nextPageToken"] = token
        data = _get("/rest/api/3/search/jql", params)
        issues.extend(data.get("issues", []))
        token = data.get("nextPageToken")
        if data.get("isLast", True) or not token:
            return issues


def last_status_in_window(key: str, since: datetime) -> str | None:
    """The last of STATUSES the ticket moved into at or after `since`."""
    latest, latest_at, start = None, None, 0
    while True:
        data = _get(f"/rest/api/3/issue/{key}/changelog", {"startAt": start, "maxResults": 100})
        for history in data.get("values", []):
            at = datetime.strptime(history["created"], "%Y-%m-%dT%H:%M:%S.%f%z")
            if at < since:
                continue
            for item in history.get("items", []):
                if item.get("field") == "status" and item.get("toString") in STATUSES:
                    if latest_at is None or at >= latest_at:
                        latest, latest_at = item["toString"], at
        start += len(data.get("values", []))
        if data.get("isLast", True) or not data.get("values"):
            return latest


def build_rows(hours: int) -> list:
    since = datetime.now(timezone.utc) - timedelta(hours=hours)
    rows = []
    for issue in find_issues(hours):
        status = last_status_in_window(issue["key"], since)
        if status:
            rows.append((issue["key"], issue["fields"].get(MGR_ORDER) or "", status))
    # Done first, then Waiting; ticket key order within each group
    rows.sort(key=lambda r: (STATUSES.index(r[2]), int(r[0].split("-")[1])))
    return rows


def write_excel(rows: list, today: datetime) -> Path:
    title = f"{TITLE} - {today:%d-%b-%Y}"
    wb = Workbook()
    ws = wb.active
    ws.title = "Hands-off"

    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:C1")

    headers = ("Ticket #", "Manager Order #", "Status")
    for col, text in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="44546A")
        cell.alignment = Alignment(horizontal="center")

    base = _get_jira_base_url()
    for r, (key, order, status) in enumerate(rows, 4):
        ws.cell(row=r, column=1, value=key).hyperlink = f"{base}/browse/{key}"
        ws.cell(row=r, column=1).font = Font(color="0563C1", underline="single")
        ws.cell(row=r, column=2, value=order)
        ws.cell(row=r, column=3, value=status)
        for col in range(1, 4):
            ws.cell(row=r, column=col).fill = FILLS[status]

    if not any(status == DONE for _, _, status in rows):
        # Right under the headers when the table is empty, else after a blank row
        note = ws.cell(row=4 + len(rows) + (1 if rows else 0), column=1, value="No shipping was done")
        note.font = Font(italic=True)

    for col, width in zip("ABC", (14, 20, 28)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A4"

    REPORT_DIR.mkdir(exist_ok=True)
    path = REPORT_DIR / f"EOD_Hands-off_Report_{today:%Y-%m-%d}.xlsx"
    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"Can't overwrite {path.name} - close it in Excel and run again.")
    return path


def send_via_outlook(path: Path, to: str, cc: str, today: datetime) -> None:
    import time
    import pywintypes
    import win32com.client
    # A cold Outlook can refuse the first connection while it is still starting up
    for attempt in range(3):
        try:
            outlook = win32com.client.Dispatch("Outlook.Application")
            break
        except pywintypes.com_error:
            if attempt == 2:
                sys.exit("Couldn't reach classic Outlook - open it once and run again with --send.")
            print("Waiting for Outlook to start...")
            time.sleep(20)
    mail = outlook.CreateItem(0)  # 0 = MailItem
    # Opening the draft makes Outlook insert the account's default signature (a
    # hidden GetInspector does too, but then Send() fails). The message goes right
    # after <body> so it sits above the signature.
    mail.Display()
    mail.To = to.replace(",", ";")
    mail.CC = cc.replace(",", ";")
    mail.Subject = f"{TITLE} - {today:%d-%b-%Y}"
    mail.HTMLBody = re.sub(r"(<body[^>]*>)", lambda m: m.group(1) + EMAIL_BODY,
                           mail.HTMLBody, count=1, flags=re.IGNORECASE)
    mail.Attachments.Add(str(path.resolve()))
    mail.Send()
    print(f"Email sent to {to}" + (f" (cc {cc})" if cc else ""))


def main():
    ap = argparse.ArgumentParser(description="Generate the EOD Hands-off Report.")
    ap.add_argument("--hours", type=int, default=8, help="look back this many hours from now (default 8)")
    ap.add_argument("--send", action="store_true", help="email the report through Outlook")
    ap.add_argument("--to", help="send only to these (comma-separated); skips the .env CC")
    ap.add_argument("--cc", help="CC these (comma-separated)")
    args = ap.parse_args()

    today = datetime.now()
    rows = build_rows(args.hours)
    path = write_excel(rows, today)

    done = sum(r[2] == DONE for r in rows)
    print(f"{len(rows)} tickets ({done} Done, {len(rows) - done} Waiting) -> {path}")

    if args.send:
        to = args.to or os.getenv("HANDOFF_EMAIL_TO", "")
        cc = args.cc or ("" if args.to else os.getenv("HANDOFF_EMAIL_CC", ""))
        if not to:
            sys.exit("No recipients: set HANDOFF_EMAIL_TO in .env or pass --to.")
        send_via_outlook(path, to, cc, today)


if __name__ == "__main__":
    main()
